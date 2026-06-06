#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

EMU_PER_INCH = 914400


def inches(value: int) -> float:
    return value / EMU_PER_INCH


def r(value: float) -> float:
    return round(value, 4)


def auto_shape_type(shape):
    try:
        return shape.auto_shape_type
    except Exception:
        return None


def text_of(shape) -> str:
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        return "\n".join(p.text for p in shape.text_frame.paragraphs).strip()
    return ""


def fill_hex(shape) -> str | None:
    try:
        rgb = shape.fill.fore_color.rgb
        if rgb is None:
            return None
        return f"{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}"
    except Exception:
        return None


def bbox(shape) -> tuple[float, float, float, float]:
    return (inches(shape.left), inches(shape.top), inches(shape.width), inches(shape.height))


def intersects(a: tuple[float, float, float, float], b: tuple[float, float, float, float], pad: float = 0.0) -> bool:
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    return ax + pad < bx + bw and bx + pad < ax + aw and ay + pad < by + bh and by + pad < ay + ah


def has_round_bullet(shape) -> bool:
    if not (getattr(shape, "has_text_frame", False) and shape.has_text_frame):
        return False
    for paragraph in shape.text_frame.paragraphs:
        ppr = paragraph._p.get_or_add_pPr()
        for child in ppr:
            if child.tag == qn("a:buChar") and child.get("char") == "●":
                return True
    return False


def bullet_runs(shape):
    if not (getattr(shape, "has_text_frame", False) and shape.has_text_frame):
        return
    for paragraph in shape.text_frame.paragraphs:
        ppr = paragraph._p.get_or_add_pPr()
        is_bullet = any(child.tag == qn("a:buChar") and child.get("char") == "●" for child in ppr)
        if is_bullet:
            for run in paragraph.runs:
                if run.text:
                    yield paragraph, run
                    break


def audit(path: Path, target_slides: int | None = None) -> dict:
    prs = Presentation(str(path))
    slide_w, slide_h = prs.slide_width, prs.slide_height
    issues: list[dict] = []
    warnings: list[dict] = []
    metrics: dict = {
        "pptx": str(path),
        "slides": len(prs.slides),
        "notes": 0,
        "media": 0,
        "right_rail_boxes": [],
        "takeaway_fonts": [],
        "rounded_radii_in": [],
        "small_pictures": [],
        "manual_short_lines": [],
    }

    with zipfile.ZipFile(path) as zf:
        bad = zf.testzip()
        metrics["zip_bad"] = bad
        metrics["slide_xml"] = len([n for n in zf.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)])
        metrics["notes_xml"] = len([n for n in zf.namelist() if re.match(r"ppt/notesSlides/notesSlide\d+\.xml$", n)])
        metrics["media"] = len([n for n in zf.namelist() if n.startswith("ppt/media/")])
        if bad:
            issues.append({"severity": "high", "slide": None, "message": f"zip test failed at {bad}"})

    if target_slides is not None and len(prs.slides) != target_slides:
        issues.append({"severity": "high", "slide": None, "message": f"slide count {len(prs.slides)} != target {target_slides}"})

    metrics["notes"] = sum(1 for s in prs.slides if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip())

    for idx, slide in enumerate(prs.slides, 1):
        pictures: list[tuple[float, float, float, float]] = []
        source_boxes: list[tuple[tuple[float, float, float, float], str]] = []
        rounded = []
        label_bars = []

        for shape in slide.shapes:
            x, y, w, h = bbox(shape)
            txt = text_of(shape)
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > slide_w + 0.02 * EMU_PER_INCH or shape.top + shape.height > slide_h + 0.02 * EMU_PER_INCH:
                issues.append({"severity": "high", "slide": idx, "message": f"shape outside bounds at {r(x)}, {r(y)}, {r(w)}, {r(h)}"})

            if txt:
                low = txt.lower()
                if any(token in low for token in ["lorem", "undefined", "nan", "xxxx"]):
                    issues.append({"severity": "high", "slide": idx, "message": f"placeholder token in text: {txt[:40]}"})
                if txt.lower().startswith("source"):
                    source_boxes.append(((x, y, w, h), txt))
                if "·" in txt:
                    issues.append({"severity": "medium", "slide": idx, "message": "hand-typed dot bullet detected"})
                for line in txt.split("\n"):
                    s = line.strip()
                    if s and not s.isdigit() and not s.lower().startswith("source") and len(s) <= 3 and w > 0.8:
                        metrics["manual_short_lines"].append({"slide": idx, "text": s, "bbox": [r(x), r(y), r(w), r(h)]})

            if shape.shape_type == 13:
                pictures.append((x, y, w, h))
                if 8 <= idx <= 19 and (w < 1.55 or h < 1.10):
                    metrics["small_pictures"].append({"slide": idx, "bbox": [r(x), r(y), r(w), r(h)]})

            if auto_shape_type(shape) == MSO_SHAPE.ROUNDED_RECTANGLE:
                rounded.append(shape)
                if hasattr(shape, "adjustments") and len(shape.adjustments):
                    metrics["rounded_radii_in"].append(r(float(shape.adjustments[0]) * min(w, h)))

            if auto_shape_type(shape) == MSO_SHAPE.RECTANGLE and fill_hex(shape) == "083578" and abs(h - 0.28) < 0.05 and x > 1.5:
                label_bars.append(shape)

            if has_round_bullet(shape) and x > 8:
                metrics["right_rail_boxes"].append({"slide": idx, "bbox": [r(x), r(y), r(w), r(h)]})
                for _paragraph, run in bullet_runs(shape):
                    if not run.text.startswith(" "):
                        issues.append({"severity": "medium", "slide": idx, "message": "bullet run missing leading space"})
                    size = run.font.size.pt if run.font.size else None
                    if size and abs(size - 18) > 0.25 and not (idx == 19 and abs(size - 17) <= 0.25):
                        issues.append({"severity": "medium", "slide": idx, "message": f"right bullet font is {size} pt"})

            if y > 6.3 and w > 8 and txt:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.text.strip() and run.font.size:
                            metrics["takeaway_fonts"].append(run.font.size.pt)

        for source_box, source_text in source_boxes:
            for picture_box in pictures:
                if intersects(source_box, picture_box, pad=0.02):
                    issues.append({"severity": "medium", "slide": idx, "message": f"source overlaps picture: {source_text[:50]}"})

        for label in label_bars:
            lx, ly, _lw, _lh = bbox(label)
            if not any(abs(bbox(frame)[0] - lx) < 0.01 and abs(bbox(frame)[1] - ly) < 0.01 for frame in rounded):
                issues.append({"severity": "medium", "slide": idx, "message": "label bar not aligned to rounded frame top-left"})

    right_pairs = {(round(item["bbox"][0], 2), round(item["bbox"][2], 2)) for item in metrics["right_rail_boxes"]}
    if len(right_pairs) > 1:
        issues.append({"severity": "medium", "slide": None, "message": f"inconsistent right rail x/width: {sorted(right_pairs)}"})

    if metrics["takeaway_fonts"] and any(abs(size - 18) > 0.25 for size in metrics["takeaway_fonts"]):
        issues.append({"severity": "medium", "slide": None, "message": f"takeaway font mismatch: {sorted(set(metrics['takeaway_fonts']))}"})

    radius_values = metrics["rounded_radii_in"]
    if radius_values and (round(min(radius_values), 2) != 0.06 or round(max(radius_values), 2) != 0.06):
        issues.append({"severity": "medium", "slide": None, "message": f"rounded radius range {min(radius_values):.4f}-{max(radius_values):.4f}"})

    if metrics["small_pictures"]:
        warnings.append({"severity": "low", "slide": None, "message": f"small pictures: {metrics['small_pictures']}"})

    return {"metrics": metrics, "issues": issues, "warnings": warnings}


def write_markdown(report: dict, path: Path) -> None:
    metrics = report["metrics"]
    high = [i for i in report["issues"] if i["severity"] == "high"]
    medium = [i for i in report["issues"] if i["severity"] == "medium"]
    low = report["warnings"]
    lines = [
        "# QA Report",
        "",
        "## Build Status",
        "",
        f"- PPTX: `{metrics['pptx']}`",
        f"- Slide count: {metrics['slides']}",
        f"- Embedded media count: {metrics['media']}",
        f"- Notes slides with text: {metrics['notes']}",
        f"- Zip test bad entry: {metrics.get('zip_bad')}",
        "",
        "## Findings",
        "",
        f"- High severity: {len(high)}",
        f"- Medium severity: {len(medium)}",
        f"- Low severity warnings: {len(low)}",
        "",
        "## Metrics",
        "",
        f"- Right rail x/width pairs: {sorted({(round(item['bbox'][0], 2), round(item['bbox'][2], 2)) for item in metrics['right_rail_boxes']})}",
        f"- Takeaway font sizes: {sorted({round(v, 1) for v in metrics['takeaway_fonts']})}",
        f"- Rounded radius range: {min(metrics['rounded_radii_in']) if metrics['rounded_radii_in'] else None} - {max(metrics['rounded_radii_in']) if metrics['rounded_radii_in'] else None}",
        f"- Small picture warnings: {len(metrics['small_pictures'])}",
        f"- Manual short-line candidates: {len(metrics['manual_short_lines'])}",
    ]
    if report["issues"]:
        lines += ["", "## Issue Details", ""]
        for issue in report["issues"]:
            lines.append(f"- {issue['severity']} / slide {issue['slide']}: {issue['message']}")
    if report["warnings"]:
        lines += ["", "## Warnings", ""]
        for warning in report["warnings"]:
            lines.append(f"- {warning['severity']}: {warning['message']}")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="Audit a v11 evidence-chain scientific defense PPTX.")
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--target-slides", type=int, default=None)
    parser.add_argument("--json", type=Path, default=None)
    parser.add_argument("--markdown", type=Path, default=None)
    args = parser.parse_args()
    report = audit(args.pptx, target_slides=args.target_slides)
    if args.json:
        args.json.parent.mkdir(parents=True, exist_ok=True)
        args.json.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    if args.markdown:
        write_markdown(report, args.markdown)
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()

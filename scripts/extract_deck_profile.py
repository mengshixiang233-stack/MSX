#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

EMU_PER_INCH = 914400


def inches(value: int) -> float:
    return round(value / EMU_PER_INCH, 4)


def auto_shape_type(shape):
    try:
        return shape.auto_shape_type
    except Exception:
        return None


def text_of(shape) -> str:
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        return "\n".join(p.text for p in shape.text_frame.paragraphs).strip()
    return ""


def fill_hex(shape) -> str | None:
    try:
        rgb = shape.fill.fore_color.rgb
        if rgb is None:
            return None
        return f"{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}"
    except Exception:
        return None


def has_round_bullet(shape) -> bool:
    if not (getattr(shape, "has_text_frame", False) and shape.has_text_frame):
        return False
    for paragraph in shape.text_frame.paragraphs:
        ppr = paragraph._p.get_or_add_pPr()
        for child in ppr:
            if child.tag == qn("a:buChar") and child.get("char") == "●":
                return True
    return False


def profile(path: Path) -> dict:
    prs = Presentation(str(path))
    out: dict = {
        "pptx": str(path),
        "slide_size_in": {"width": inches(prs.slide_width), "height": inches(prs.slide_height)},
        "slides": [],
    }
    with zipfile.ZipFile(path) as zf:
        out["zip_test"] = zf.testzip()
        out["slide_xml_count"] = len([n for n in zf.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)])
        out["notes_xml_count"] = len([n for n in zf.namelist() if re.match(r"ppt/notesSlides/notesSlide\d+\.xml$", n)])
        out["media_count"] = len([n for n in zf.namelist() if n.startswith("ppt/media/")])

    for idx, slide in enumerate(prs.slides, 1):
        slide_info = {
            "slide": idx,
            "texts": [],
            "pictures": [],
            "rounded_rectangles": [],
            "label_bars": [],
            "right_bullet_boxes": [],
        }
        for shape in slide.shapes:
            x, y, w, h = [inches(v) for v in (shape.left, shape.top, shape.width, shape.height)]
            bbox = {"x": x, "y": y, "w": w, "h": h}
            txt = text_of(shape)
            if txt:
                slide_info["texts"].append({**bbox, "text": txt[:500]})
            if shape.shape_type == 13:
                slide_info["pictures"].append(bbox)
            if auto_shape_type(shape) == MSO_SHAPE.ROUNDED_RECTANGLE:
                radius = None
                if hasattr(shape, "adjustments") and len(shape.adjustments):
                    radius = round(float(shape.adjustments[0]) * min(w, h), 4)
                slide_info["rounded_rectangles"].append({**bbox, "radius_in": radius, "fill": fill_hex(shape)})
            if auto_shape_type(shape) == MSO_SHAPE.RECTANGLE:
                if fill_hex(shape) == "083578" and abs(h - 0.28) < 0.05 and x > 1.5:
                    slide_info["label_bars"].append(bbox)
            if has_round_bullet(shape) and x > 8:
                slide_info["right_bullet_boxes"].append({**bbox, "text": txt[:300]})
        out["slides"].append(slide_info)
    return out


def main() -> None:
    parser = argparse.ArgumentParser(description="Extract PPTX layout profile for evidence-chain defense decks.")
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--json", type=Path, default=None)
    args = parser.parse_args()
    data = profile(args.pptx)
    text = json.dumps(data, ensure_ascii=False, indent=2)
    if args.json:
        args.json.parent.mkdir(parents=True, exist_ok=True)
        args.json.write_text(text + "\n", encoding="utf-8")
    print(text)


if __name__ == "__main__":
    main()
